"""Tests de extract_runs: cobertura estrictamente ≥ python-docx."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document


def _normalize(text: str) -> str:
    """Colapsa whitespace para comparar contenido real, no separadores."""
    return re.sub(r"\s+", " ", text).strip()

from app.io.docx_extract import count_coverage, extract_runs

FIXTURES = Path(__file__).resolve().parent.parent / "ejemplos"
DOCX_FILES = sorted(p for p in FIXTURES.glob("*.docx") if not p.name.startswith("~$"))


@pytest.mark.parametrize("docx_path", DOCX_FILES, ids=lambda p: p.name)
def test_extract_runs_basico(docx_path: Path) -> None:
    doc = extract_runs(docx_path)
    assert doc.runs, f"No se extrajo ningún run de {docx_path.name}"
    assert doc.full_text.strip(), "full_text vacío"
    # Los offsets de cada run deben apuntar dentro de full_text y coincidir.
    for run in doc.runs:
        assert doc.full_text[run.char_start:run.char_end] == run.text


@pytest.mark.parametrize("docx_path", DOCX_FILES, ids=lambda p: p.name)
def test_cobertura_mayor_o_igual_que_python_docx(docx_path: Path) -> None:
    """Garantía dura: extract_runs ve al menos tanto texto como python-docx.

    python-docx sólo recorre doc.paragraphs del cuerpo principal. Nuestro
    extractor además lee headers, footers, tablas, footnotes. Por lo tanto
    el número de caracteres del nuevo extractor debe ser >= al del viejo.
    """
    d_old = Document(str(docx_path))
    old_text = _normalize("\n".join(p.text for p in d_old.paragraphs))

    new_doc = extract_runs(docx_path)
    new_text = _normalize(new_doc.full_text)

    # Contenido: cada carácter no-whitespace del viejo debe estar presente
    # (en la misma cantidad o más) en el nuevo. Comparamos por longitud
    # normalizada, que es robusta a diferencias de separadores.
    assert len(new_text) >= len(old_text), (
        f"Regresión de contenido en {docx_path.name}: "
        f"nuevo={len(new_text)} chars normalizados, viejo={len(old_text)}"
    )
    # Y adicionalmente: todo el texto viejo debe aparecer como subsecuencia
    # de palabras en el nuevo. Verificamos con una heurística simple:
    # cada palabra larga del viejo (>3 chars) debe estar en el nuevo.
    old_words = {w for w in old_text.split() if len(w) > 3}
    missing = [w for w in old_words if w not in new_text]
    assert not missing, f"Palabras perdidas en {docx_path.name}: {missing[:5]}"


@pytest.mark.parametrize("docx_path", DOCX_FILES, ids=lambda p: p.name)
def test_metricas_de_cobertura(docx_path: Path) -> None:
    doc = extract_runs(docx_path)
    cov = count_coverage(doc)
    assert cov["num_runs"] > 0
    assert cov["num_chars"] > 0
    # Al menos el document.xml debe estar presente.
    assert "word/document.xml" in cov["parts"]
